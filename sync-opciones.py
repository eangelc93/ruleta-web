"""
Lee "Opciones de la ruleta.xlsx" y genera opciones.json.
Uso:  python sync-opciones.py
Luego sube el cambio a GitHub para que Vercel lo publique.
"""
from __future__ import annotations

import json
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
XLSX = ROOT / "Opciones de la ruleta.xlsx"
OUT = ROOT / "opciones.json"


def plain(value: str) -> str:
    text = unicodedata.normalize("NFD", str(value or "").strip().lower())
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def main() -> None:
    if not XLSX.exists():
        raise SystemExit(f"No está el Excel: {XLSX.name}")

    wb = load_workbook(XLSX, data_only=True)
    ws = wb["Opciones"]
    headers = [str(cell.value or "").strip() for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    cols = {plain(name): index for index, name in enumerate(headers)}
    if "nombre" not in cols:
        raise SystemExit('La hoja Opciones necesita la columna "Nombre".')

    opciones = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        raw_name = row[cols["nombre"]]
        if raw_name is None:
            continue
        nombre = str(raw_name).strip()
        if not nombre:
            continue

        def get(*aliases: str) -> str:
            for alias in aliases:
                if alias in cols and row[cols[alias]] not in (None, ""):
                    return str(row[cols[alias]]).strip()
            return ""

        item = {"nombre": nombre}
        rueda = get("texto en la ruleta", "rueda", "texto")
        detalle = get("detalle", "informacion", "descripcion")
        color = get("color")
        icono = get("icono")
        peso = get("peso")
        activa = get("activa")
        if rueda:
            item["rueda"] = rueda
        if detalle:
            item["detalle"] = detalle
        if color:
            item["color"] = color
        if icono:
            item["icono"] = icono
        if peso:
            try:
                item["peso"] = int(float(peso))
            except ValueError:
                pass
        if plain(activa) in {"no", "n", "0", "false", "inactiva"}:
            item["activa"] = False
        opciones.append(item)

    info_map = {
        "titulo linea 1": "titulo1",
        "titulo linea 2": "titulo2",
        "frase": "frase",
        "encabezado": "encabezado",
        "paso 1 titulo": "paso1Titulo",
        "paso 1 texto": "paso1Texto",
        "paso 2 titulo": "paso2Titulo",
        "paso 2 texto": "paso2Texto",
        "paso 3 titulo": "paso3Titulo",
        "paso 3 texto": "paso3Texto",
        "etiqueta del resultado": "etiquetaResultado",
        "mensaje del resultado": "mensajeResultado",
        "boton siguiente": "botonSiguiente",
        "segundos para el siguiente": "segundos",
    }
    informacion = {}
    if "Informacion" in wb.sheetnames:
        for row in wb["Informacion"].iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            label = plain(str(row[0]))
            if label == "campo" or label not in info_map:
                continue
            key = info_map[label]
            value = row[1]
            if key == "segundos":
                try:
                    informacion[key] = int(float(value))
                except (TypeError, ValueError):
                    pass
            elif value is not None:
                informacion[key] = str(value)

    pack = {"opciones": opciones, "informacion": informacion}
    OUT.write_text(json.dumps(pack, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Listo: {OUT.name} ({len(opciones)} opciones)")


if __name__ == "__main__":
    main()
